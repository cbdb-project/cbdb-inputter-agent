"""Turn a "CBDB 人物資料標準化 template" workbook into a staging batch.

    python -m cbdb_agent from-workbook --xlsx <workbook> --case cases/<id>/workbook.yaml

The template is a standing format, not one job's spreadsheet: a 年譜 or biography is
broken down sentence by sentence into one sheet per CBDB table (基本資料 → BIOG_MAIN,
官名 → POSTED_TO_OFFICE_DATA, 社會關係 → ASSOC_DATA, …), every row already coded and
already compared with CBDB. Its own 說明 sheet is the specification; COLUMNS below is
the part of it this reader relies on, and `check_template()` refuses a workbook whose
說明 maps one of those columns to a different database field.

What each row's `disambiguation` turns into:

* **new** - a `create`. A person the template could not find in CBDB carries a
  temporary id (`TMP-001`); that becomes a `basicinformation` create with
  `person_id: NEW`, every row of theirs names it as its person, and a relative's or
  associate's `c_kin_id`/`c_assoc_id` is a `{"ref": ...}` to it.
* **update** - an `update` of exactly the fields the template lists after 「可補：」
  in 待確認事項, which is its record of "CBDB has this row and leaves these empty".
  Nothing else on the row is written: an existing CBDB value is never overwritten
  from here. A year field brings its 年號/年號年/範圍 companions with it. Which
  existing row to update is the case file's `update_targets` - the template says
  "CBDB has this", not which composite key - and a missing target is a finding, not a
  guess.
* **exist** - nothing. Where the template says the row differs from CBDB
  (「與 CBDB 既有記錄不同」) the difference is a finding for the reviewer.

A row the template marks 待確認, or whose note asks a human to judge (「請判斷」…),
gets a conflict that has to be resolved before `submit`, and so does every row naming
a person whose own 基本資料 row is 待確認. Approving or deferring is the reviewer's;
this module does not re-judge the coder's reading.

The batch's source book is the case file's `source_text`: an existing `textid`, or a
`create` that the batch makes first and every row cites through `{"ref": ...}`. A
著述 row whose text is 需新建 does the same with its own title. Both creates are
global reference data with no delete path (AGENTS.md rule 12), so each carries a
conflict and goes nowhere without an explicit approval in review.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any

import yaml

from .mirror_check import is_pair_update
from .mirror_check import new_conflict as mirror_conflict
from .models import find_spec_by_alias
from .staging import Conflict, ConflictOption, Proposal, StagingBatch

# sheet -> (resource alias, id prefix). The prefix keeps proposal ids ASCII; the
# template's own 記錄編號 is quoted at the head of every source_quote.
SHEETS: dict[str, tuple[str, str]] = {
    "基本資料": ("basicinformation", "bio"),
    "地址": ("addresses", "addr"),
    "別名": ("altnames", "alt"),
    "著述": ("texts", "text"),
    "官名": ("postings", "post"),
    "入仕": ("entries", "entry"),
    "社會區分": ("statuses", "status"),
    "親屬": ("kinship", "kin"),
    "社會關係": ("associations", "assoc"),
    "社交機構": ("social_institutions", "inst"),
}

_YEAR_FIELDS = {
    "起始年": ("c_firstyear", "c_fy_nh_code", "c_fy_nh_year", "c_fy_range", "c_fy_month"),
    "結束年": ("c_lastyear", "c_ly_nh_code", "c_ly_nh_year", "c_ly_range", "c_ly_month"),
}


def _span(prefix: str, fields: tuple[str, ...], month: bool = True) -> dict[str, str]:
    year, nh, nh_year, rng, mon = fields
    cols = {f"{prefix}(西元)": year, f"{prefix}號代碼": nh, f"{prefix}號年": nh_year,
            f"{prefix}範圍代碼": rng}
    if month and mon:
        cols[f"{prefix[:-1]}月" if prefix.endswith("年") else f"{prefix}月"] = mon
    return cols


_SOURCE_COLS = {"出處文本ID": "c_source", "頁碼": "c_pages", "備註": "c_notes"}

# column label -> database field, per sheet. Only columns that are written; the
# Chinese-term twin of each code column is for human reading and never sent.
COLUMNS: dict[str, dict[str, str]] = {
    "基本資料": {
        "人物姓名": "c_name_chn", "姓名拼音": "c_name", "姓氏": "c_surname_chn",
        "名": "c_mingzi_chn", "性別代碼": "c_female", "朝代代碼": "c_dy",
        "生年(西元)": "c_birthyear", "生年年號代碼": "c_by_nh_code",
        "生年年號年": "c_by_nh_year", "生年月": "c_by_month", "生年日": "c_by_day",
        "生年範圍代碼": "c_by_range",
        "卒年(西元)": "c_deathyear", "卒年年號代碼": "c_dy_nh_code",
        "卒年年號年": "c_dy_nh_year", "卒年月": "c_dy_month", "卒年日": "c_dy_day",
        "卒年範圍代碼": "c_dy_range",
        "享年": "c_death_age", "享年範圍代碼": "c_death_age_range",
        "指數年": "c_index_year", "指數年類型代碼": "c_index_year_type_code",
        "備註": "c_notes",
        # 出處文本ID/頁碼 are not BIOG_MAIN columns: they become a `sources` row.
    },
    "地址": {
        "地址類型代碼": "c_addr_type", "地址ID": "c_addr_id", "序號": "c_sequence",
        **_span("起始年", _YEAR_FIELDS["起始年"]), **_span("結束年", _YEAR_FIELDS["結束年"]),
        **_SOURCE_COLS,
    },
    "別名": {
        "別名(標準)": "c_alt_name_chn", "別名拼音": "c_alt_name",
        "別名類型代碼": "c_alt_name_type_code", "序號": "c_sequence", **_SOURCE_COLS,
    },
    "著述": {
        "文本ID": "c_textid", "角色代碼": "c_role_id", **_SOURCE_COLS,
        # 年(西元)/年號代碼/年號年/年範圍代碼 are BIOG_TEXT_DATA columns the server
        # does not accept (models.py's texts comment) - reported, never sent.
    },
    "官名": {
        "官名ID": "c_office_id", "官名朝代代碼": "c_dy", "任命類型代碼": "c_appt_code",
        "就任狀況代碼": "c_assume_office_code", "任官地ID": "c_addr",
        "官職類別代碼": "c_office_category_id", "序號": "c_sequence",
        **_span("起始年", _YEAR_FIELDS["起始年"]), **_span("結束年", _YEAR_FIELDS["結束年"]),
        **_SOURCE_COLS,
    },
    "入仕": {
        "入仕途徑代碼": "c_entry_code", "考試等第": "c_exam_rank", "序號": "c_sequence",
        "年(西元)": "c_year", "年號代碼": "c_entry_nh_id", "年號年": "c_entry_nh_year",
        "年範圍代碼": "c_entry_range", "年齡": "c_age", **_SOURCE_COLS,
    },
    "社會區分": {
        "社會區分代碼": "c_status_code", "補充說明": "c_supplement", "序號": "c_sequence",
        **_span("起始年", _YEAR_FIELDS["起始年"], month=False),
        **_span("結束年", _YEAR_FIELDS["結束年"], month=False),
        **_SOURCE_COLS,
    },
    "親屬": {"親屬人物ID": "c_kin_id", "親屬關係代碼": "c_kin_code", **_SOURCE_COLS},
    "社會關係": {
        "關係人物ID": "c_assoc_id", "關係類型代碼": "c_assoc_code",
        "相關文本篇名": "c_text_title", "場合代碼": "c_occasion_code",
        "關係地點ID": "c_addr_id", "序號": "c_sequence",
        "起始年(西元)": "c_assoc_first_year", "起始年號代碼": "c_assoc_fy_nh_code",
        "起始年號年": "c_assoc_fy_nh_year", "起始月": "c_assoc_fy_month",
        "起始年範圍代碼": "c_assoc_fy_range",
        "結束年(西元)": "c_assoc_last_year", "結束年號代碼": "c_assoc_ly_nh_code",
        "結束年號年": "c_assoc_ly_nh_year", "結束月": "c_assoc_ly_month",
        "結束年範圍代碼": "c_assoc_ly_range",
        **_SOURCE_COLS,
        # 文類代碼 is ASSOC_DATA.c_litgenre_code, which the associations whitelist
        # does not accept (a 422, not a silent drop) - reported, never sent.
    },
    "社交機構": {
        "機構名稱代碼": "c_inst_name_code", "機構代碼": "c_inst_code",
        "人物角色代碼": "c_bi_role_code",
        "起始年(西元)": "c_bi_begin_year", "起始年號代碼": "c_bi_by_nh_code",
        "起始年號年": "c_bi_by_nh_year", "起始年範圍代碼": "c_bi_by_range",
        "結束年(西元)": "c_bi_end_year", "結束年號代碼": "c_bi_ey_nh_code",
        "結束年號年": "c_bi_ey_nh_year", "結束年範圍代碼": "c_bi_ey_range",
        **_SOURCE_COLS,
    },
}

# Where the template names the table column and the API takes a pseudo-field for it:
# a posting's place is POSTED_TO_ADDR_DATA.c_addr_id, sent as postings' `c_addr` list.
_DECLARED_AS = {("官名", "任官地ID"): "c_addr_id"}

# Template columns that name a real database column the API will not take.
UNSENDABLE: dict[str, dict[str, str]] = {
    "著述": {"年(西元)": "c_year"},
    "社會關係": {"文類代碼": "c_litgenre_code"},
}

# A 可補 label that is a year brings the rest of its year with it: filling c_firstyear
# and leaving its 年號 empty would be half a date.
_YEAR_COMPANIONS = {
    "起始年(西元)": ("起始年號代碼", "起始年號年", "起始年範圍代碼"),
    "結束年(西元)": ("結束年號代碼", "結束年號年", "結束年範圍代碼"),
    "年(西元)": ("年號代碼", "年號年", "年範圍代碼"),
    "生年(西元)": ("生年年號代碼", "生年年號年", "生年範圍代碼"),
    "卒年(西元)": ("卒年年號代碼", "卒年年號年", "卒年範圍代碼"),
}

# Codes the template stores as text and CBDB stores as text: kept as strings, so
# the leading zero survives.
_TEXT_CODES_COLUMNS = {"c_index_year_type_code", "c_text_type_id"}

# Non-key sequence columns. On a person CBDB already has, the template's own
# numbering (1, 2, 3 within this workbook) would interleave with the existing rows'
# and mean nothing; it is left to the server. Where c_sequence is part of the key
# (addresses, entries, statuses) it stays.
_ORDER_ONLY_SEQUENCE = {"postings", "altnames", "associations"}

_REVIEW_MARKERS = ("請判斷", "請確認", "請核對", "請人工")
_IDENTITY_NOTE = re.compile(r"人物身分見\s*基本-\d+")
_FILL_LIST = re.compile(r"可補[:：]\s*([^　\n；;]+)")
_DIFFERS = re.compile(r"與 CBDB 既有記錄不同[:：]\s*([^　\n]+)")


class WorkbookError(ValueError):
    pass


@dataclass
class CaseConfig:
    batch_id: str
    source_text: dict[str, Any]
    update_targets: dict[str, dict[str, Any]] = field(default_factory=dict)
    new_texts: dict[str, dict[str, Any]] = field(default_factory=dict)
    source_excerpt: str | None = None

    @classmethod
    def load(cls, path: str) -> "CaseConfig":
        with open(path, encoding="utf-8") as f:
            raw = yaml.safe_load(f) or {}
        source = raw.get("source_text") or {}
        if ("textid" in source) == ("create" in source):
            raise WorkbookError(
                f"{path}: source_text needs exactly one of `textid` (an existing "
                "TEXT_CODES row) or `create` (a title this batch creates)")
        return cls(batch_id=raw["batch_id"], source_text=source,
                   update_targets=raw.get("update_targets") or {},
                   new_texts=raw.get("new_texts") or {},
                   source_excerpt=raw.get("source_excerpt"))


@dataclass
class Row:
    sheet: str
    record: str
    values: dict[str, Any]

    def get(self, column: str) -> Any:
        return _clean(self.values.get(column))


@dataclass
class Result:
    batch: StagingBatch
    findings: list[str]


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


# --- reading -------------------------------------------------------------------


def read_workbook(path: str) -> dict[str, list[Row]]:
    try:
        import openpyxl
    except ImportError as exc:   # a reader-only dependency, see requirements-dev.txt
        raise WorkbookError("reading a workbook needs openpyxl "
                            "(pip install -r requirements-dev.txt)") from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        check_template(wb)
        sheets: dict[str, list[Row]] = {}
        for name in SHEETS:
            if name not in wb.sheetnames:
                raise WorkbookError(f"{path}: no sheet {name!r} - not this template")
            rows = list(wb[name].iter_rows(values_only=True))
            header = [str(h).strip() if h is not None else None for h in rows[0]]
            known = (set(COLUMNS[name]) | set(UNSENDABLE.get(name, {}))
                     | {"記錄編號", "人物ID", "disambiguation", "標準化狀態", "待確認事項"})
            missing = sorted(c for c in known if c not in header)
            if missing:
                raise WorkbookError(f"sheet {name!r} has no column(s) {missing}")
            sheets[name] = [
                Row(name, str(r[header.index("記錄編號")]).strip(), dict(zip(header, r)))
                for r in rows[1:] if any(v is not None for v in r)
            ]
        return sheets
    finally:
        wb.close()


def check_template(wb) -> None:
    """Refuse a workbook whose 說明 maps a column COLUMNS relies on elsewhere.

    The 欄位對照表 there has (sheet, column, database field) per row. A template
    revision that renamed a column fails `read_workbook` on the header; one that
    re-pointed a column at a different field would otherwise be read silently wrong.
    """
    if "說明" not in wb.sheetnames:
        raise WorkbookError("no 說明 sheet - not a 人物資料標準化 template workbook")
    declared: dict[tuple[str, str], str] = {}
    for row in wb["說明"].iter_rows(values_only=True):
        if len(row) >= 3 and row[0] in SHEETS and row[1] and row[2]:
            target = str(row[2]).strip()
            declared[(row[0], str(row[1]).strip())] = target.split(".")[-1]
    for sheet, columns in COLUMNS.items():
        for column, db_field in columns.items():
            said = declared.get((sheet, column))
            expected = _DECLARED_AS.get((sheet, column), db_field)
            if said is not None and said not in ("—", expected):
                raise WorkbookError(
                    f"說明 maps {sheet}.{column} to {said!r}, this reader to "
                    f"{db_field!r} - the template has changed; update COLUMNS")


# --- the conversion ------------------------------------------------------------


def build_batch(sheets: dict[str, list[Row]], case: CaseConfig) -> Result:
    findings: list[str] = []
    proposals: list[Proposal] = []
    people = {r.get("人物ID"): r for r in sheets["基本資料"]}
    tmp_ids = {pid: _person_proposal_id(pid) for pid in people if _is_tmp(pid)}

    def person_ref(pid: Any) -> Any:
        if _is_tmp(pid):
            if pid not in tmp_ids:
                raise WorkbookError(f"{pid} is referenced but has no 基本資料 row")
            return {"ref": tmp_ids[pid]}
        return int(pid)

    # The batch's own source book.
    if "create" in case.source_text:
        source: Any = {"ref": "txt-source"}
        proposals.append(_text_create(
            "txt-source", _guard_text_codes(case.source_text["create"]),
            case.source_text.get("evidence", ""),
            "the source of every row in this batch"))
    else:
        source = int(case.source_text["textid"])

    texts_by_record: dict[str, Any] = {}
    for row in sheets["著述"]:
        if row.get("文本ID狀態") == "需新建" and row.get("disambiguation") == "new":
            pid = f"txt-{_record_slug(row.record)}"
            extra = dict(case.new_texts.get(row.record) or {})
            if "textid" in extra:
                # Created since the workbook was written: cite it, never create it
                # again - TEXT_CODES has no delete.
                texts_by_record[row.record] = int(extra["textid"])
                continue
            evidence = extra.pop("evidence", "")
            create = {"c_title_chn": row.get("標準篇名")}
            if row.get("文本類型代碼"):
                create["c_text_type_id"] = _typed("c_text_type_id", row.get("文本類型代碼"))
            create.update(_guard_text_codes(extra))
            proposals.append(_text_create(pid, create, evidence, f"{row.record}"))
            texts_by_record[row.record] = {"ref": pid}

    # People first: each later row's person_id may name one of these.
    # An existing person whose own row is open - or refused (不入庫/無法對應: the
    # id may not be the person meant) - makes every row naming them a question.
    identity_open = {_person_key(pid): r for pid, r in people.items()
                     if r.get("標準化狀態") in ("待確認", "不入庫", "無法對應")}
    for row in sheets["基本資料"]:
        proposals.extend(_person_rows(row, source, tmp_ids, findings))

    for sheet, (resource, _prefix) in SHEETS.items():
        if sheet == "基本資料":
            continue
        for row in sheets[sheet]:
            proposal = _sub_row(row, resource, source, person_ref, tmp_ids,
                                texts_by_record, case, identity_open, findings)
            if proposal is not None:
                proposals.append(proposal)

    findings.extend(_unsendable_findings(sheets))
    batch = StagingBatch(
        batch_id=case.batch_id,
        source_excerpt=case.source_excerpt,
        proposals=proposals,
        batch_notes=_batch_notes(findings),
    )
    return Result(batch=batch, findings=findings)


def _person_rows(row: Row, source: Any, tmp_ids: dict, findings: list[str]) -> list[Proposal]:
    pid, kind = row.get("人物ID"), row.get("disambiguation")
    status = row.get("標準化狀態")
    if status in ("不入庫", "無法對應"):
        findings.append(f"{row.record}: {status} - not proposed "
                        f"({_strip_identity(row.get('待確認事項')) or 'no reason given'})")
        if _is_tmp(pid):
            # Rows that name this person must fail loudly, not be emitted.
            tmp_ids.pop(pid, None)
        return []
    if kind == "exist":
        _note_differences(row, findings)
        return []
    if kind == "update" and _is_tmp(pid):
        raise WorkbookError(f"{row.record}: a TMP- person cannot be an update")
    quote = _quote(row)
    out: list[Proposal] = []
    if kind == "new":
        if not _is_tmp(pid):
            raise WorkbookError(f"{row.record}: a new person must carry a TMP- id, not {pid!r}")
        changes = _mapped(row, "基本資料")
        changes.pop("c_notes", None)
        notes = clean_notes(row.get("備註"))
        if notes:
            changes["c_notes"] = notes
        own_id = tmp_ids[pid]
        out.append(Proposal(
            id=own_id, resource="basicinformation", operation="create",
            person_id="NEW", target_pk=None, changes=changes, source_quote=quote,
            confidence=_confidence(row),
            conflicts=_row_conflicts(own_id, row, {})))
        out.append(_source_row(f"{own_id}-src", own_id, source, row, main=True))
    elif kind == "update":
        _note_differences(row, findings)
        fill = _fill_fields(row, "基本資料", findings)
        if not fill:
            findings.append(f"{row.record}: marked update with nothing listed after 可補 - "
                            "nothing proposed")
            return []
        own_id = f"bio-{pid}"
        out.append(Proposal(
            id=own_id, resource="basicinformation", operation="update",
            person_id=int(pid), target_pk=None, changes=fill, source_quote=quote,
            confidence=_confidence(row),
            conflicts=_row_conflicts(own_id, row, {})))
        out.append(_source_row(f"{own_id}-src", int(pid), source, row, main=False))
    return out


def _source_row(pid: str, person: Any, source: Any, row: Row, *, main: bool) -> Proposal:
    target_pk: dict[str, Any] = {"c_textid": source}
    if row.get("頁碼"):
        pages = row.get("頁碼")
        if isinstance(pages, float) and pages.is_integer():
            pages = int(pages)
        target_pk["c_pages"] = str(pages)
    return Proposal(
        id=pid, resource="sources", operation="create", person_id=person,
        target_pk=target_pk, changes={"c_main_source": 1} if main else {},
        source_quote=f"[{row.record}] the source book this batch cites",
        confidence="high")


def _sub_row(row: Row, resource: str, source: Any, person_ref, tmp_ids: dict,
             texts_by_record: dict, case: CaseConfig, identity_open: dict,
             findings: list[str]) -> Proposal | None:
    kind, status = row.get("disambiguation"), row.get("標準化狀態")
    if status in ("不入庫", "無法對應"):
        findings.append(f"{row.record}: {status} - not proposed "
                        f"({_strip_identity(row.get('待確認事項')) or 'no reason given'})")
        return None
    if kind == "exist":
        _note_differences(row, findings)
        return None
    pid_cell = row.get("人物ID")
    if _is_tmp(pid_cell) and pid_cell not in tmp_ids:
        raise WorkbookError(f"{row.record} belongs to {pid_cell}, who has no 基本資料 "
                            "row to create (or is marked 不入庫/無法對應)")
    person_id = tmp_ids[pid_cell] if _is_tmp(pid_cell) else int(pid_cell)
    spec = find_spec_by_alias(resource)
    proposal_id = f"{SHEETS[row.sheet][1]}-{_record_slug(row.record)}"
    referenced = [pid_cell] + [row.get(c) for c in ("親屬人物ID", "關係人物ID") if row.get(c)]
    identity = {_person_key(p): identity_open[_person_key(p)] for p in referenced
                if not _is_tmp(p) and _person_key(p) in identity_open}

    if kind == "update":
        _note_differences(row, findings)
        changes = _fill_fields(row, row.sheet, findings)
        unsendable = [c for c in UNSENDABLE.get(row.sheet, {}) if c in _fill_labels(row)]
        if unsendable:
            findings.append(f"{row.record}: 可補 {'、'.join(unsendable)}, which the "
                            f"{resource} endpoint does not accept - not proposed")
        if not changes:
            if not unsendable:
                findings.append(f"{row.record}: marked update with nothing listed after "
                                f"可補 - nothing proposed")
            return None
        target = case.update_targets.get(row.record)
        if not target:
            findings.append(f"{row.record}: update has no target in the case file "
                            f"(which existing {resource} row?) - not proposed")
            return None
        conflicts = _row_conflicts(proposal_id, row, identity)
        if is_pair_update(resource, "update"):
            conflicts.append(mirror_conflict(proposal_id, resource))
        return Proposal(
            id=proposal_id, resource=resource, operation="update",
            person_id=person_id, target_pk=dict(target), changes=changes,
            source_quote=_quote(row), confidence=_confidence(row),
            conflicts=conflicts)

    if kind != "new":
        raise WorkbookError(f"{row.record}: disambiguation {kind!r} is not new/update/exist")

    changes = _mapped(row, row.sheet)
    own = changes.get("c_source")
    if own is not None and own != source:
        findings.append(f"{row.record}: the row cites 出處文本ID {own}; the batch's "
                        f"source replaced it")
    changes["c_source"] = source
    notes = clean_notes(row.get("備註"))
    changes.pop("c_notes", None)
    if notes:
        changes["c_notes"] = notes
    if resource == "postings" and "c_addr" in changes:
        changes["c_addr"] = [int(changes["c_addr"])]
    if resource == "kinship":
        changes["c_kin_id"] = person_ref(changes["c_kin_id"])
    if resource == "associations":
        changes["c_assoc_id"] = person_ref(changes["c_assoc_id"])
        # Required, non-empty key columns with documented sentinels (docs/04 s7).
        changes.setdefault("c_text_title", "[n/a]")
        changes.setdefault("c_assoc_first_year", -9999)
        for f in ("c_kin_code", "c_kin_id", "c_assoc_kin_code", "c_assoc_kin_id"):
            changes.setdefault(f, 0)
    if resource == "texts":
        if row.record in texts_by_record:
            changes["c_textid"] = texts_by_record[row.record]
        elif "c_textid" not in changes:
            findings.append(f"{row.record}: no 文本ID and not 需新建 - not proposed")
            return None
    if resource in _ORDER_ONLY_SEQUENCE and not _is_tmp(pid_cell):
        changes.pop("c_sequence", None)

    target_pk = {f: changes[f] for f in spec.pk_fields
                 if f in changes and f != "c_personid"
                 and f not in spec.server_assigned_pk_fields}
    return Proposal(
        id=proposal_id, resource=resource, operation="create",
        person_id=person_id, target_pk=target_pk, changes=changes,
        source_quote=_quote(row), confidence=_confidence(row),
        conflicts=_row_conflicts(proposal_id, row, identity))


def _text_create(pid: str, create: dict, evidence: str, used_for: str) -> Proposal:
    title = create.get("c_title_chn")
    return Proposal(
        id=pid, resource="text-codes", operation="create", person_id=0,
        target_pk={}, changes=dict(create),
        source_quote=f"New TEXT_CODES title, cited by {used_for}. {evidence}".strip(),
        confidence="medium",
        conflicts=[Conflict(
            id=f"{pid}-c1", field="c_title_chn",
            description=(
                f"Create {title!r} in TEXT_CODES. This is global reference data with "
                "no delete path, and c_title_chn cannot be changed afterwards "
                "(AGENTS.md rule 12). Deferring it holds back every row that cites it."),
            options=[
                ConflictOption(value="confirmed", rationale="create the title"),
                ConflictOption(value="defer", rationale="do not create it in this batch"),
            ],
            agent_suggestion=None,
            agent_reasoning="creating global reference data is the reviewer's call",
        )],
    )


# --- row helpers ---------------------------------------------------------------


def _mapped(row: Row, sheet: str) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for column, db_field in COLUMNS[sheet].items():
        value = row.get(column)
        if value is None:
            continue
        out[db_field] = _typed(db_field, value)
    return out


def _guard_text_codes(values: dict) -> dict:
    """The case file's own TEXT_CODES values. YAML reads an unquoted `010110` as
    the octal integer 4168, so a code column that is not a string is refused here
    exactly as it is from the workbook."""
    return {k: (_typed(k, v) if k in _TEXT_CODES_COLUMNS else v) for k, v in values.items()}


def _typed(db_field: str, value: Any) -> Any:
    if db_field in _TEXT_CODES_COLUMNS:
        if not isinstance(value, str):
            # A number-typed cell has lost its leading zeros ("01" -> 1), and the
            # padding cannot be guessed. TEXT_CODES rows cannot be corrected later.
            raise WorkbookError(
                f"{db_field} must be a text cell in the workbook, got {value!r}")
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"-?\d+", value) and db_field not in (
            "c_pages", "c_notes", "c_name_chn", "c_alt_name_chn", "c_exam_rank",
            "c_text_title", "c_supplement"):
        return int(value)
    return value


def _fill_labels(row: Row) -> list[str]:
    match = _FILL_LIST.search(row.get("待確認事項") or "")
    if not match:
        return []
    return [label.strip() for label in re.split(r"[、,，]", match.group(1)) if label.strip()]


def _fill_fields(row: Row, sheet: str, findings: list[str] | None = None) -> dict[str, Any]:
    """The 可補 fields and their year companions, valued from the row.

    A listed label that names no sendable column, or whose cell is empty, is a
    finding - silently writing less than the template asked for is how a gap goes
    unnoticed. (A label the API refuses is reported by the caller.)
    """
    out: dict[str, Any] = {}
    for label in _fill_labels(row):
        if label in UNSENDABLE.get(sheet, {}):
            continue
        if label not in COLUMNS[sheet] or row.get(label) is None:
            if findings is not None:
                why = "not a column" if label not in COLUMNS[sheet] else "empty"
                findings.append(f"{row.record}: 可補 lists {label!r}, which is {why} "
                                f"- not written")
            continue
        for part in (label, *_YEAR_COMPANIONS.get(label, ())):
            db_field = COLUMNS[sheet].get(part)
            value = row.get(part)
            if db_field is None or value is None:
                continue
            out[db_field] = (clean_notes(value) if db_field == "c_notes"
                             else _typed(db_field, value))
    return {k: v for k, v in out.items() if v is not None}


def _row_conflicts(pid: str, row: Row, identity: dict[str, Row]) -> list[Conflict]:
    """A conflict for whatever the workbook leaves open on this row itself, plus one
    per open person it names.

    Two of the template's notes are settled by how this module works, not open: the
    可補 list is what an update writes, and 「與 CBDB 既有記錄不同」 is a difference in
    a field that is never overwritten (it becomes a finding). A row that is 待確認
    for those reasons alone has nothing left for a reviewer to decide.
    """
    conflicts: list[Conflict] = []
    raw = _strip_identity(row.get("待確認事項"))
    note = _strip_identity(row.get("待確認事項"), settled=True)
    only_settled = raw is not None and note is None
    if (row.get("標準化狀態") == "待確認" and not only_settled) or any(
            m in (note or "") for m in _REVIEW_MARKERS):
        conflicts.append(Conflict(
            id=f"{pid}-c1", field="row",
            description=f"{row.record} is open in the workbook: {note or '(no note)'}",
            options=[
                ConflictOption(value="confirmed", rationale="write this row as coded"),
                ConflictOption(value="defer", rationale="leave it out of this batch"),
            ],
        ))
    for n, (person, person_row) in enumerate(sorted(identity.items(), key=str), 1):
        conflicts.append(Conflict(
            id=f"{pid}-id{n}", field="person",
            description=(f"Who is {person} ({person_row.get('人物姓名')})? "
                         f"{person_row.record} is open: "
                         f"{person_row.get('待確認事項') or '(no note)'}"),
            options=[
                ConflictOption(value="confirmed",
                               rationale=f"{person} is the person the source means"),
                ConflictOption(value="defer", rationale="leave this row out"),
            ],
        ))
    return conflicts


def _strip_identity(text: Any, *, settled: bool = False) -> str | None:
    """The note without its 「人物身分見 基本-nnn」 pointers (identity is asked once,
    from the person's own row) and, with `settled`, without the 可補 / differs notes
    `_row_conflicts` explains."""
    if not text:
        return None
    parts = [p.strip() for p in str(text).split("　")]
    kept = [p for p in parts if p and not _IDENTITY_NOTE.fullmatch(p)
            and not (settled and (_FILL_LIST.search(p) or _DIFFERS.search(p)))]
    return "　".join(kept) or None


def _note_differences(row: Row, findings: list[str]) -> None:
    match = _DIFFERS.search(row.get("待確認事項") or "")
    if match:
        findings.append(f"{row.record} (CBDB kept): workbook differs from the existing "
                        f"row - {match.group(1).strip()}")


def _quote(row: Row) -> str:
    text = row.get("原文") or row.get("原文依據") or ""
    return f"[{row.record}] {text}".strip()


def _confidence(row: Row) -> str:
    return "high" if row.get("標準化狀態") == "已標準化" else "medium"


def _unsendable_findings(sheets: dict[str, list[Row]]) -> list[str]:
    out = []
    for sheet, columns in UNSENDABLE.items():
        for column, db_field in columns.items():
            hit = [r.record for r in sheets[sheet]
                   if r.get(column) is not None and r.get("disambiguation") == "new"]
            if hit:
                out.append(f"{column} ({db_field}) is not accepted by the API and was "
                           f"left out of: {', '.join(hit)}")
    return out


def _batch_notes(findings: list[str]) -> str:
    head = ("Generated from a 人物資料標準化 workbook by cbdb_agent.person_workbook. "
            "Rows the workbook marks `exist` are not proposed. Findings:")
    return "\n".join([head, *[f"- {f}" for f in findings]]) if findings else head


# --- c_notes -------------------------------------------------------------------

# The template's merge step leaves its own bookkeeping in 備註: which source rows of
# the intermediate sheet a row was merged from, and by which key. That is a fact
# about the workbook, meaningless to anyone reading CBDB, and it would be written
# into a public row. Removed here; everything else is the coder's prose and stays.
_MERGE_SEGMENT = re.compile(r"^【併入來源列[^】]*】|本列由來源列")
_SOURCE_ROW_PAREN = re.compile(r"[（(]來源列\s*[\d;；、,，\-–~～\s]+[）)]")


def clean_notes(text: Any) -> str | None:
    if not text:
        return None
    parts = [p.strip() for p in unicodedata.normalize("NFC", str(text)).split("　")]
    kept = [_SOURCE_ROW_PAREN.sub("", p) for p in parts
            if p and not _MERGE_SEGMENT.search(p)]
    out = "　".join(p for p in kept if p)
    return out or None


# --- ids -----------------------------------------------------------------------


def _person_key(pid: Any) -> Any:
    """An existing person's id as an int, whether the cell held 845 or "845"."""
    if _is_tmp(pid):
        return pid
    try:
        return int(str(pid).strip())
    except (TypeError, ValueError):
        return pid


def _is_tmp(pid: Any) -> bool:
    return isinstance(pid, str) and pid.upper().startswith("TMP-")


def _person_proposal_id(pid: str) -> str:
    return f"p-{pid.lower()}"


def _record_slug(record: str) -> str:
    """官名-002 -> 002; the sheet is already in the proposal id's prefix."""
    return record.split("-", 1)[1] if "-" in record else record
